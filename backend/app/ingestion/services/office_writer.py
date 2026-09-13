"""Create bounded editable Office payloads from Library text."""

from __future__ import annotations

from io import BytesIO

from app.core.errors import ApiError


def text_to_docx(text: str) -> bytes:
    try:
        from docx import Document
    except Exception as exc:  # pragma: no cover
        raise ApiError(
            code="DOCX_EXTRACTOR_DEPENDENCY_MISSING",
            message="DOCX support is unavailable.",
            status_code=503,
        ) from exc
    document = Document()
    for line in text.splitlines() or [""]:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def text_to_pptx(text: str) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as exc:  # pragma: no cover
        raise ApiError(
            code="PPTX_EXTRACTOR_DEPENDENCY_MISSING",
            message="PPTX support is unavailable.",
            status_code=503,
        ) from exc
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    lines = text.splitlines() or [""]
    for offset in range(0, len(lines), 40):
        slide = presentation.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12.1), Inches(6.2))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = "\n".join(lines[offset : offset + 40])
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(16)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def text_to_xlsx(text: str) -> bytes:
    """Build a workbook from the extractor's stable ``Sheet/R/C`` text form."""
    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover
        raise ApiError(
            code="XLSX_EXTRACTOR_DEPENDENCY_MISSING",
            message="XLSX support is unavailable.",
            status_code=503,
        ) from exc
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = None
    for line in text.splitlines() or [""]:
        if line.startswith("## Sheet:"):
            sheet = workbook.create_sheet(line.removeprefix("## Sheet:").strip()[:31] or "Sheet")
            continue
        if sheet is None:
            sheet = workbook.create_sheet("Sheet")
        if line.startswith("R") and ":" in line:
            row_label, values = line.split(":", 1)
            try:
                row = int(row_label[1:])
            except ValueError:
                row = sheet.max_row + 1
            for value in values.split(" | "):
                if "=" not in value or not value.startswith("C"):
                    continue
                column_label, cell_value = value.split("=", 1)
                try:
                    column = int(column_label[1:])
                except ValueError:
                    continue
                sheet.cell(row=row, column=column, value=cell_value)
        elif line:
            sheet.cell(row=sheet.max_row + 1, column=1, value=line)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
